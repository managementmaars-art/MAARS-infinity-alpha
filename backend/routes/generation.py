"""File serving, document/image/video generation endpoints."""
import uuid
import io
import asyncio
import base64
import logging
import mimetypes
from pathlib import Path
from fastapi import APIRouter, HTTPException, Depends, Request
from fastapi.responses import StreamingResponse
from db import db
from auth import get_current_user, User
from shared.constants import UPLOAD_DIR
from shared.utils import get_api_keys

logger = logging.getLogger(__name__)
router = APIRouter()

@router.get("/files/{filename}")
async def serve_file(filename: str):
    """Serve a generated file"""
    file_path = UPLOAD_DIR / filename
    if not file_path.exists():
        raise HTTPException(status_code=404, detail="File not found")
    
    content_type = mimetypes.guess_type(str(file_path))[0] or "application/octet-stream"
    
    def file_iter():
        with open(file_path, "rb") as f:
            while chunk := f.read(8192):
                yield chunk
    
    return StreamingResponse(
        file_iter(),
        media_type=content_type,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )

@router.post("/generate/document")
async def generate_document(request: Request, current_user: User = Depends(get_current_user)):
    """Generate PDF, Excel, DOCX, CSV, or TXT documents"""
    body = await request.json()
    doc_type = body.get("type", "pdf").lower()  # pdf, xlsx, docx, csv, txt
    title = body.get("title", "Generated Document")
    content = body.get("content", "")
    rows = body.get("rows", [])  # For spreadsheets: list of lists
    
    if not content and not rows:
        raise HTTPException(status_code=400, detail="Content or rows required")
    
    file_id = uuid.uuid4().hex[:10]
    
    try:
        if doc_type == "pdf":
            from reportlab.lib.pagesizes import letter
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
            from reportlab.lib.styles import getSampleStyleSheet
            
            filename = f"{file_id}_{title.replace(' ', '_')[:30]}.pdf"
            filepath = UPLOAD_DIR / filename
            
            doc = SimpleDocTemplate(str(filepath), pagesize=letter)
            styles = getSampleStyleSheet()
            story = [
                Paragraph(title, styles['Title']),
                Spacer(1, 20),
            ]
            for para in content.split("\n"):
                if para.strip():
                    story.append(Paragraph(para.strip(), styles['Normal']))
                    story.append(Spacer(1, 8))
            doc.build(story)
        
        elif doc_type == "xlsx":
            from openpyxl import Workbook
            from openpyxl.styles import Font
            
            filename = f"{file_id}_{title.replace(' ', '_')[:30]}.xlsx"
            filepath = UPLOAD_DIR / filename
            
            wb = Workbook()
            ws = wb.active
            ws.title = title[:31]
            
            if rows:
                for r_idx, row in enumerate(rows, 1):
                    for c_idx, val in enumerate(row, 1):
                        cell = ws.cell(row=r_idx, column=c_idx, value=val)
                        if r_idx == 1:
                            cell.font = Font(bold=True)
            else:
                ws.cell(row=1, column=1, value=title).font = Font(bold=True)
                for i, line in enumerate(content.split("\n"), 2):
                    ws.cell(row=i, column=1, value=line.strip())
            
            wb.save(str(filepath))
        
        elif doc_type == "docx":
            from docx import Document
            from docx.shared import Pt
            
            filename = f"{file_id}_{title.replace(' ', '_')[:30]}.docx"
            filepath = UPLOAD_DIR / filename
            
            doc = Document()
            doc.add_heading(title, level=1)
            for para in content.split("\n"):
                if para.strip():
                    doc.add_paragraph(para.strip())
            doc.save(str(filepath))
        
        elif doc_type == "csv":
            import csv as csv_module
            
            filename = f"{file_id}_{title.replace(' ', '_')[:30]}.csv"
            filepath = UPLOAD_DIR / filename
            
            with open(filepath, "w", newline="") as f:
                writer = csv_module.writer(f)
                if rows:
                    for row in rows:
                        writer.writerow(row)
                else:
                    for line in content.split("\n"):
                        writer.writerow([line.strip()])
        
        elif doc_type == "txt":
            filename = f"{file_id}_{title.replace(' ', '_')[:30]}.txt"
            filepath = UPLOAD_DIR / filename
            
            with open(filepath, "w") as f:
                f.write(content)
        
        else:
            raise HTTPException(status_code=400, detail=f"Unsupported doc type: {doc_type}. Use: pdf, xlsx, docx, csv, txt")
        
        return {
            "filename": filename,
            "url": f"/files/{filename}",
            "type": doc_type,
            "title": title,
            "size": filepath.stat().st_size
        }
    
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Document generation error: {e}")
        raise HTTPException(status_code=500, detail=f"Document generation failed: {str(e)}")

@router.post("/generate/image")
async def generate_image(request: Request, current_user: User = Depends(get_current_user)):
    """Generate an image through the MAARS media router (Gemini/OpenAI)."""
    body = await request.json()
    prompt = body.get("prompt", "")
    quality = body.get("quality", "standard")  # "standard" | "premium"
    size = body.get("size", "1024x1024")

    if not prompt:
        raise HTTPException(status_code=400, detail="Prompt is required")

    try:
        from services.media_router import route_image
        from services.billing.media_billing import bill_and_run
        # Estimate: standard tier image = 20 credits (gpt-image-1/$0.02), premium = 40 (dall-e-3/$0.04)
        est_model = "dall-e-3" if quality == "premium" else "gpt-image-1"
        image_bytes, meta, billing = await bill_and_run(
            current_user.user_id, "image", est_model, 1,
            route_image, prompt=prompt, quality=quality, size=size,
        )

        file_id = uuid.uuid4().hex[:10]
        filename = f"{file_id}_generated.png"
        filepath = UPLOAD_DIR / filename
        with open(filepath, "wb") as f:
            f.write(image_bytes)

        image_b64 = base64.b64encode(image_bytes).decode()
        # Scrub infrastructure details — clients see only the MAARS brand
        # and whether enhancement happened. Provider/model/cost_usd are
        # stripped so clients never know which backend produced the image.
        from shared.response_scrubber import scrub_meta, scrub
        return {
            "filename": filename,
            "url": f"/files/{filename}",
            "preview": f"data:image/png;base64,{image_b64}",
            "prompt": prompt,
            "router": scrub_meta(meta),
            "billing": scrub(billing),
        }
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Image generation error: {e}")
        raise HTTPException(status_code=500, detail=f"Image generation failed: {str(e)}")


# Video gating — Sora-2 at $0.40 per 4-sec clip = 400 credits. Free/Starter/Essential
# don't have enough budget for a single video even in the best case, so gate it cleanly
# with a 402 that tells the user which plan they need rather than a confusing "insufficient
# credits" after a partial flow.
_VIDEO_MIN_PLANS = {"basic", "standard", "professional", "advanced", "business",
                    "agency", "studio", "enterprise", "corporate", "elite", "owner"}


async def _get_user_plan_id(user_id: str) -> str:
    """Return the caller's current plan_id (defaults to 'free'). Admin → 'owner'."""
    from db import db as _db
    from auth import ADMIN_EMAIL
    user = await _db.users.find_one({"user_id": user_id}, {"_id": 0, "email": 1})
    if user and user.get("email") == ADMIN_EMAIL:
        return "owner"
    sub = await _db.subscriptions.find_one({"user_id": user_id}, {"_id": 0, "plan_id": 1})
    return (sub or {}).get("plan_id", "free")


@router.post("/generate/video")
async def generate_video(request: Request, current_user: User = Depends(get_current_user)):
    """Generate a video through the MAARS media router (Sora-2 today)."""
    body = await request.json()
    prompt = body.get("prompt", "")
    size = body.get("size", "1280x720")
    duration = body.get("duration", 4)

    if not prompt:
        raise HTTPException(status_code=400, detail="Prompt is required")
    if size not in ("1280x720", "1792x1024", "1024x1792", "1024x1024"):
        size = "1280x720"
    if duration not in (4, 8, 12):
        duration = 4

    # Plan gate — Sora-2 is too expensive for the cheapest plans.
    plan_id = await _get_user_plan_id(current_user.user_id)
    if plan_id not in _VIDEO_MIN_PLANS:
        raise HTTPException(status_code=402, detail={
            "error": {
                "message": (f"Video generation requires the Basic plan or higher. "
                            f"Your current plan ({plan_id}) doesn't include video. "
                            f"A single 4-second video costs 400 credits; Basic includes 1,200 credits/mo."),
                "type": "plan_upgrade_required",
                "code": "video_plan_gate",
                "current_plan": plan_id,
                "required_plan": "basic",
                "video_cost_credits": 400 * (duration // 4),
            }
        })

    try:
        from services.media_router import route_video
        from services.billing.media_billing import bill_and_run
        video_bytes, meta, billing = await bill_and_run(
            current_user.user_id, "video", "sora-2", duration,
            route_video, prompt=prompt, duration=duration, size=size,
        )

        file_id = uuid.uuid4().hex[:10]
        filename = f"{file_id}_video.mp4"
        filepath = UPLOAD_DIR / filename
        with open(filepath, "wb") as f:
            f.write(video_bytes)

        from shared.response_scrubber import scrub_meta, scrub
        return {
            "filename": filename,
            "url": f"/files/{filename}",
            "size": size,
            "duration": duration,
            "prompt": prompt,
            "router": scrub_meta(meta),
            "billing": scrub(billing),
        }
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Video generation error: {e}")
        raise HTTPException(status_code=500, detail=f"Video generation failed: {str(e)}")

